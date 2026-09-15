import io
import hashlib
import hmac
import json
import os
import re
import secrets
import smtplib
import sqlite3
import unicodedata
import uuid
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import parse_qs
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape as xml_escape

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
TEMPLATE_PATH = BASE_DIR / "templates" / "Tagesrapport App.xlsx"
DATA_DIR = Path(os.environ.get("DATA_DIR", BASE_DIR / "data"))
DB_PATH = DATA_DIR / "tagesrapporte.db"
MAX_BODY = 2 * 1024 * 1024
PROJECT_LEADERS = [
    "Imer Abazi", "Fidan Salihi", "Damiano Giarrusso", "Megzon Morina",
    "Sandro Sargenti", "Armando Baldacci", "Nico Gavoldi", "Remo Brunner",
]


def utc_now():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def db_connection():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS reports (
            id TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Entwurf',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            payload TEXT NOT NULL
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS employees (
            id TEXT PRIMARY KEY,
            name TEXT NOT NULL UNIQUE COLLATE NOCASE,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS project_leaders (
            id TEXT PRIMARY KEY,
            name TEXT NOT NULL UNIQUE COLLATE NOCASE,
            email TEXT NOT NULL UNIQUE COLLATE NOCASE,
            password_hash TEXT,
            password_salt TEXT,
            updated_at TEXT NOT NULL
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS login_sessions (
            token_hash TEXT PRIMARY KEY,
            project_leader_id TEXT NOT NULL,
            expires_at TEXT NOT NULL,
            FOREIGN KEY(project_leader_id) REFERENCES project_leaders(id)
        )
        """
    )
    for name in PROJECT_LEADERS:
        connection.execute(
            "INSERT OR IGNORE INTO project_leaders (id, name, email, updated_at) VALUES (?, ?, ?, ?)",
            (str(uuid.uuid4()), name, leader_email(name), utc_now()),
        )
    connection.commit()
    return connection


def ascii_slug(value):
    normalized = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", normalized.lower()).strip("-") or "unbekannt"


def leader_email(name):
    parts = ascii_slug(name).split("-")
    local = f"{parts[0]}.{parts[-1]}" if len(parts) > 1 else parts[0]
    return f"{local}@polygongroup.com"


def hash_password(password, salt=None):
    salt_bytes = bytes.fromhex(salt) if salt else secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt_bytes, 310_000)
    return digest.hex(), salt_bytes.hex()


def project_leader_session(environ):
    header = environ.get("HTTP_AUTHORIZATION", "")
    if not header.startswith("Bearer "):
        return None
    token = header[7:].strip()
    if not token:
        return None
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    with db_connection() as connection:
        row = connection.execute(
            """SELECT p.id, p.name, p.email, s.expires_at
               FROM login_sessions s JOIN project_leaders p ON p.id = s.project_leader_id
               WHERE s.token_hash = ?""",
            (token_hash,),
        ).fetchone()
        if not row:
            return None
        if datetime.fromisoformat(row["expires_at"]) <= datetime.now(timezone.utc):
            connection.execute("DELETE FROM login_sessions WHERE token_hash = ?", (token_hash,))
            connection.commit()
            return None
    return dict(row)


def auth_context(environ):
    leader = project_leader_session(environ)
    if leader:
        return {"kind": "leader", **leader}
    expected = os.environ.get("APP_PASSWORD", "").strip()
    provided = environ.get("HTTP_X_APP_KEY", "")
    if (expected and hmac.compare_digest(provided, expected)) or not expected:
        return {"kind": "admin"}
    return None


def json_response(start_response, payload, status="200 OK", headers=None):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    final_headers = [("Content-Type", "application/json; charset=utf-8"), ("Content-Length", str(len(body)))]
    final_headers.extend(headers or [])
    start_response(status, final_headers)
    return [body]


def text_response(start_response, body, content_type, status="200 OK", headers=None):
    encoded = body if isinstance(body, bytes) else body.encode("utf-8")
    final_headers = [("Content-Type", content_type), ("Content-Length", str(len(encoded)))]
    final_headers.extend(headers or [])
    start_response(status, final_headers)
    return [encoded]


def read_json(environ):
    try:
        length = int(environ.get("CONTENT_LENGTH") or 0)
    except ValueError:
        length = 0
    if length <= 0 or length > MAX_BODY:
        raise ValueError("Ungültige Anfragegrösse")
    raw = environ["wsgi.input"].read(length)
    return json.loads(raw.decode("utf-8"))


def safe_title(data):
    location = (data.get("location") or {}).get("name") or "Tagesrapport"
    date = data.get("date") or datetime.now().date().isoformat()
    return f"{location} – {date}"[:160]


def normalize_report(data):
    if not isinstance(data, dict):
        raise ValueError("Rapportdaten fehlen")
    required = ["date", "reporterName", "region", "projectLead", "damageType"]
    missing = [field for field in required if not str(data.get(field, "")).strip()]
    location = data.get("location") or {}
    if not str(location.get("name", "")).strip():
        missing.append("Projektname / Projektnummer")
    if missing:
        raise ValueError("Bitte Pflichtfelder ausfüllen: " + ", ".join(missing))
    data["areas"] = [area for area in data.get("areas", []) if area in {"fh", "gb", "tr"}]
    return data


def clear_and_mark(ws, choices, selected, text_cell=None):
    for label, cell in choices.items():
        ws[cell] = "x" if label == selected else None
    if text_cell and selected not in choices:
        ws[text_cell] = selected


def write_rows(ws, rows, start_row, columns, max_rows):
    for offset in range(max_rows):
        row_number = start_row + offset
        item = rows[offset] if offset < len(rows) else {}
        for key, column in columns.items():
            value = item.get(key, "")
            ws[f"{column}{row_number}"] = value if value not in (None, "") else None


def build_excel(report):
    workbook = load_workbook(TEMPLATE_PATH)
    ws = workbook[workbook.sheetnames[0]]

    ws["A1"] = f"Tagesrapport – {report.get('date', '')}"
    ws["A2"] = f"Ausgefüllt von: {report.get('reporterName', '')}"
    clear_and_mark(ws, {
        "Bern / Wallis": "C9", "Mittelland / Nordwest": "C10", "Tessin": "C11", "Zürich / Ostschweiz": "C12"
    }, report.get("region"))
    clear_and_mark(ws, {
        "Imer Abazi": "C15", "Fidan Salihi": "C16", "Damiano Giarrusso": "C17", "Megzon Morina": "C18",
        "Sandro Sargenti": "C19", "Armando Baldacci": "C20", "Nico Gavoldi": "C21", "Remo Brunner": "C22"
    }, report.get("projectLead"))
    damage = report.get("damageType", "")
    clear_and_mark(ws, {"Feuer": "C24", "Wasser": "C25"}, damage)
    ws["C26"] = "x" if damage not in {"Feuer", "Wasser"} else None
    ws["E26"] = damage if damage not in {"Feuer", "Wasser"} else None

    location = report.get("location") or {}
    ws["E28"], ws["E29"], ws["E30"] = location.get("name"), location.get("address"), location.get("city")
    ws["E31"] = f"Datum: {report.get('date', '')}"

    areas = set(report.get("areas") or [])
    for cell, key in [("B37", "fh"), ("C37", "gb"), ("E37", "tr")]:
        ws[cell] = "x" if key in areas else None

    work = report.get("work") or {}
    write_rows(ws, work.get("fh", []), 68, {"role": "A", "name": "B", "hours": "D", "temp": "E"}, 13)
    write_rows(ws, work.get("gb", []), 68, {"role": "F", "name": "G", "hours": "I", "temp": "J"}, 13)
    write_rows(ws, work.get("tr", []), 68, {"role": "K", "name": "L", "hours": "N", "temp": "O"}, 13)

    extras = report.get("extras") or {}
    for key, col in [("fh", "D"), ("gb", "I"), ("tr", "N")]:
        values = extras.get(key) or {}
        ws[f"{col}85"] = values.get("saturday") or None
        ws[f"{col}86"] = values.get("sunday") or None
        if key != "tr":
            ws[f"{col}91"] = "x" if values.get("leadFlat") == "Telefonisch" else None
            ws[f"{col}92"] = "x" if values.get("leadFlat") == "Vor Ort" else None

    vehicles = report.get("vehicles") or {}
    for key, cols in [("fh", ("B", "C", "D")), ("gb", ("G", "H", "I")), ("tr", ("L", "M", "N"))]:
        by_type = {str(item.get("type")): item for item in vehicles.get(key, [])}
        for row, vehicle_type, aliases in [(103, "PW", ["PW"]), (104, "Bus", ["Bus", "Transporter"]), (105, "Kein Fahrzeug", ["Kein Fahrzeug", "Keine (San. Im Lager)"])]:
            item = next((by_type.get(alias) for alias in aliases if by_type.get(alias)), {})
            label_col = chr(ord(cols[0]) - 1)
            ws[f"{label_col}{row}"] = vehicle_type
            ws[f"{cols[0]}{row}"] = item.get("count") or None
            ws[f"{cols[1]}{row}"] = item.get("billing") or None
            ws[f"{cols[2]}{row}"] = item.get("amount") or None

    equipment = report.get("equipment") or {}
    write_rows(ws, equipment.get("fh", []), 131, {"item": "A", "hours": "B"}, 10)
    write_rows(ws, equipment.get("gb", []), 131, {"item": "F", "hours": "G"}, 10)
    write_rows(ws, equipment.get("tr", []), 131, {"item": "K", "built": "L", "removed": "M"}, 10)

    products = report.get("products") or {}
    write_rows(ws, products.get("fh", []), 161, {"item": "A", "amount": "B"}, 10)
    write_rows(ws, products.get("gb", []), 161, {"item": "F", "amount": "G"}, 10)
    write_rows(ws, products.get("tr", []), 161, {"item": "K", "amount": "L"}, 10)

    consumables = report.get("consumables") or {}
    write_rows(ws, consumables.get("fh", []), 192, {"item": "A", "amount": "B"}, 6)
    write_rows(ws, consumables.get("gb", []), 192, {"item": "F", "amount": "G"}, 6)
    write_rows(ws, consumables.get("tr", []), 192, {"item": "K", "amount": "L"}, 6)

    disposal = report.get("disposal") or {}
    disposal_rows = {"Entsorgung vor Ort in Mulde": 223, "Entsorgung Lager Polygon": 225, "Entsorgung bei lokalem Entsorgungszentrum": 227}
    for key, col in [("fh", "C"), ("gb", "H")]:
        entry = disposal.get(key, "")
        chosen = entry.get("method", "") if isinstance(entry, dict) else entry
        for label, row in disposal_rows.items():
            ws[f"{col}{row}"] = "x" if chosen == label else None
        if isinstance(entry, dict) and chosen:
            amount_col = "D" if key == "fh" else "I"
            unit_col = "E" if key == "fh" else "J"
            chosen_row = disposal_rows.get(chosen, 223)
            ws[f"{amount_col}{chosen_row}"] = entry.get("amount") or None
            ws[f"{unit_col}{chosen_row}"] = entry.get("unit") or None

    notes = str(report.get("notes") or "").strip()
    if notes:
        ws["A230"] = "Bemerkungen"
        ws.merge_cells("A231:O236")
        ws["A231"] = notes
        ws["A231"].alignment = ws["A231"].alignment.copy(wrap_text=True, vertical="top")

    ws.sheet_view.showGridLines = False
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.read()


def pdf_table(rows, widths=None):
    table = Table(rows, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#082A3C")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D8E2E7")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F8")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def build_pdf(report):
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="RapportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, textColor=colors.HexColor("#082A3C"), alignment=TA_LEFT, spaceAfter=12))
    styles.add(ParagraphStyle(name="Section", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor("#0D6C7F"), spaceBefore=12, spaceAfter=6))
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=15*mm, bottomMargin=15*mm, title=safe_title(report))
    story = [Paragraph("Tagesrapport", styles["RapportTitle"])]
    location = report.get("location") or {}
    basics = [
        ["Datum", report.get("date", "")], ["Ausgefüllt von", report.get("reporterName", "")],
        ["Region", report.get("region", "")], ["Projektleiter", report.get("projectLead", "")],
        ["Schadenart", report.get("damageType", "")], ["Projektname / Projektnummer", location.get("name", "")],
        ["Adresse", location.get("address", "")], ["Ort", location.get("city", "")],
        ["Bereiche", ", ".join({"fh": "Fahrhabe", "gb": "Gebäude", "tr": "Trocknung"}.get(x, x) for x in (report.get("areas") or []))],
    ]
    story.append(pdf_table([["Angabe", "Wert"]] + basics, [48*mm, 125*mm]))

    area_names = {"fh": "Fahrhabe", "gb": "Gebäude", "tr": "Trocknung"}
    for key in report.get("areas") or []:
        story.extend([Paragraph(area_names.get(key, key), styles["Section"]), Paragraph("Arbeitsstunden", styles["Heading3"])])
        work_rows = [["Funktion / Tätigkeit", "Personaltyp", "Name", "Stunden", "Temporärbüro"]]
        for item in (report.get("work") or {}).get(key, []):
            if item.get("name") or item.get("hours"):
                work_rows.append([item.get("role", ""), item.get("employeeType", ""), item.get("name", ""), item.get("hours", ""), item.get("temp", "")])
        if len(work_rows) > 1:
            story.append(pdf_table(work_rows, [38*mm, 25*mm, 42*mm, 18*mm, 45*mm]))
        extra = (report.get("extras") or {}).get(key) or {}
        extra_rows = [["Zuschlag / Pauschale", "Wert"]]
        if extra.get("saturday") not in ("", None, 0):
            extra_rows.append(["Samstag (Std.)", extra.get("saturday")])
        if extra.get("sunday") not in ("", None, 0):
            extra_rows.append(["Sonntag (Std.)", extra.get("sunday")])
        if extra.get("leadFlat"):
            extra_rows.append(["Projektleiter-Pauschale", extra.get("leadFlat")])
        if len(extra_rows) > 1:
            story.append(Spacer(1, 2*mm))
            story.append(pdf_table(extra_rows, [70*mm, 98*mm]))

        for section_key, title, columns in [
            ("vehicles", "Fahrzeuge", [("type", "Fahrzeug"), ("count", "Anzahl"), ("billing", "Abrechnung"), ("amount", "P / KM")]),
            ("equipment", "Geräte", [("item", "Gerät"), ("hours", "Stunden"), ("built", "Aufgebaut"), ("removed", "Abgebaut")]),
            ("products", "Sanierungsprodukte / Spezialgeräte", [("item", "Position"), ("amount", "Menge")]),
            ("consumables", "Verbrauchsmaterial", [("item", "Position"), ("amount", "Menge")]),
        ]:
            items = [item for item in (report.get(section_key) or {}).get(key, []) if any(v not in ("", None, 0) for v in item.values())]
            if items:
                story.append(Paragraph(title, styles["Heading3"]))
                story.append(pdf_table([[label for _, label in columns]] + [[item.get(field, "") for field, _ in columns] for item in items]))
        disposal = (report.get("disposal") or {}).get(key) or {}
        if isinstance(disposal, dict) and disposal.get("method"):
            story.append(Paragraph("Entsorgung", styles["Heading3"]))
            story.append(pdf_table([["Entsorgungsweg", "Menge", "Einheit"], [disposal.get("method", ""), disposal.get("amount", ""), disposal.get("unit", "")]]))
        story.append(Spacer(1, 4*mm))

    notes = str(report.get("notes") or "").strip()
    if notes:
        story.extend([Paragraph("Bemerkungen", styles["Section"]), Paragraph(xml_escape(notes).replace("\n", "<br/>"), styles["BodyText"])])
    story.extend([Spacer(1, 8*mm), Paragraph(f"Status: {report.get('status', 'Entwurf')}", styles["BodyText"])])
    document.build(story)
    return output.getvalue()


def save_pdf_to_folder(report):
    today = datetime.now(ZoneInfo("Europe/Zurich")).date().isoformat()
    relative = Path("project_leaders") / ascii_slug(report.get("projectLead")) / today / f"rapport-{report['id']}.pdf"
    destination = DATA_DIR / relative
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(build_pdf(report))
    return str(relative)


def public_base_url(environ):
    configured = os.environ.get("BASE_URL", "").strip().rstrip("/")
    if configured:
        return configured
    scheme = (environ.get("HTTP_X_FORWARDED_PROTO") or environ.get("wsgi.url_scheme") or "https").split(",")[0].strip()
    host = (environ.get("HTTP_X_FORWARDED_HOST") or environ.get("HTTP_HOST") or "localhost").split(",")[0].strip()
    if scheme not in {"http", "https"} or not re.fullmatch(r"[A-Za-z0-9.:-]+", host):
        raise ValueError("BASE_URL ist nicht korrekt konfiguriert")
    return f"{scheme}://{host}"


def send_project_leader_email(report, environ):
    host = os.environ.get("SMTP_HOST", "").strip()
    username = os.environ.get("SMTP_USERNAME", "").strip()
    password = os.environ.get("SMTP_PASSWORD", "")
    sender = os.environ.get("SMTP_FROM", username).strip()
    if not all([host, username, password, sender]):
        return False, "SMTP ist nicht vollständig konfiguriert"
    port = int(os.environ.get("SMTP_PORT", "587"))
    recipient = leader_email(report.get("projectLead", ""))
    reporter = report.get("reporterName") or "Eine Person"
    link = f"{public_base_url(environ)}/?portal=1&report={report['id']}"
    message = EmailMessage()
    message["Subject"] = f"Neuer Tagesrapport: {safe_title(report)}"
    message["From"] = sender
    message["To"] = recipient
    message.set_content(f"{reporter} hat einen Rapport ausgefüllt.\n\nRapport öffnen und prüfen:\n{link}\n")
    try:
        if port == 465:
            with smtplib.SMTP_SSL(host, port, timeout=20) as smtp:
                smtp.login(username, password)
                smtp.send_message(message)
        else:
            with smtplib.SMTP(host, port, timeout=20) as smtp:
                smtp.ehlo()
                smtp.starttls()
                smtp.ehlo()
                smtp.login(username, password)
                smtp.send_message(message)
        return True, recipient
    except Exception as error:
        print(f"Email error for report {report.get('id')}: {error}", flush=True)
        return False, "E-Mail konnte nicht gesendet werden"


def serve_static(start_response, path):
    mapping = {
        "/": (STATIC_DIR / "index.html", "text/html; charset=utf-8"),
        "/styles.css": (STATIC_DIR / "styles.css", "text/css; charset=utf-8"),
        "/app.js": (STATIC_DIR / "app.js", "application/javascript; charset=utf-8"),
    }
    item = mapping.get(path)
    if not item or not item[0].exists():
        return text_response(start_response, "Nicht gefunden", "text/plain; charset=utf-8", "404 Not Found")
    return text_response(start_response, item[0].read_bytes(), item[1])


def application(environ, start_response):
    method = environ.get("REQUEST_METHOD", "GET").upper()
    path = environ.get("PATH_INFO", "/")

    if path == "/health":
        return json_response(start_response, {"status": "ok"})
    if path == "/api/config":
        return json_response(start_response, {
            "authRequired": bool(os.environ.get("APP_PASSWORD", "").strip()),
            "projectLeaders": [{"name": name, "email": leader_email(name)} for name in PROJECT_LEADERS],
            "languages": ["de", "fr", "it"],
        })

    try:
        if path == "/api/auth/login" and method == "POST":
            data = read_json(environ)
            name = str(data.get("name") or "").strip()
            password = str(data.get("password") or "")
            with db_connection() as connection:
                leader = connection.execute("SELECT * FROM project_leaders WHERE name = ? COLLATE NOCASE", (name,)).fetchone()
                if not leader or not leader["password_hash"] or not leader["password_salt"]:
                    return json_response(start_response, {"error": "Login nicht eingerichtet oder ungültig"}, "401 Unauthorized")
                digest, _ = hash_password(password, leader["password_salt"])
                if not hmac.compare_digest(digest, leader["password_hash"]):
                    return json_response(start_response, {"error": "Login nicht eingerichtet oder ungültig"}, "401 Unauthorized")
                token = secrets.token_urlsafe(36)
                expires = datetime.now(timezone.utc) + timedelta(hours=12)
                connection.execute("DELETE FROM login_sessions WHERE expires_at <= ?", (utc_now(),))
                connection.execute(
                    "INSERT INTO login_sessions (token_hash, project_leader_id, expires_at) VALUES (?, ?, ?)",
                    (hashlib.sha256(token.encode("utf-8")).hexdigest(), leader["id"], expires.replace(microsecond=0).isoformat()),
                )
                connection.commit()
            return json_response(start_response, {"token": token, "projectLead": leader["name"], "expiresAt": expires.isoformat()})

        context = auth_context(environ)
        if path.startswith("/api/") and not context:
            return json_response(start_response, {"error": "Zugriffsschlüssel oder Login ungültig"}, "401 Unauthorized")

        if path == "/api/project-leaders" and method == "GET":
            if context["kind"] != "admin":
                return json_response(start_response, {"error": "Keine Berechtigung"}, "403 Forbidden")
            with db_connection() as connection:
                rows = connection.execute("SELECT id, name, email, password_hash FROM project_leaders ORDER BY name").fetchall()
            return json_response(start_response, {"projectLeaders": [{"id": row["id"], "name": row["name"], "email": row["email"], "loginConfigured": bool(row["password_hash"])} for row in rows]})

        if path == "/api/project-leaders/password" and method == "POST":
            if context["kind"] != "admin":
                return json_response(start_response, {"error": "Keine Berechtigung"}, "403 Forbidden")
            data = read_json(environ)
            name = str(data.get("name") or "").strip()
            password = str(data.get("password") or "")
            if len(password) < 10:
                raise ValueError("Das Passwort muss mindestens 10 Zeichen lang sein")
            digest, salt = hash_password(password)
            with db_connection() as connection:
                result = connection.execute("UPDATE project_leaders SET password_hash = ?, password_salt = ?, updated_at = ? WHERE name = ? COLLATE NOCASE", (digest, salt, utc_now(), name))
                connection.execute("DELETE FROM login_sessions WHERE project_leader_id IN (SELECT id FROM project_leaders WHERE name = ? COLLATE NOCASE)", (name,))
                connection.commit()
            if not result.rowcount:
                raise ValueError("Projektleiter nicht gefunden")
            return json_response(start_response, {"updated": True, "name": name})

        if path == "/api/reports" and method == "GET":
            with db_connection() as connection:
                rows = connection.execute("SELECT id, title, status, created_at, updated_at, payload FROM reports ORDER BY updated_at DESC LIMIT 250").fetchall()
            reports = []
            for row in rows:
                item = dict(row)
                payload = json.loads(item.pop("payload"))
                if context["kind"] == "leader" and payload.get("projectLead") != context["name"]:
                    continue
                item["projectLead"] = payload.get("projectLead", "Ohne Zuteilung")
                item["date"] = payload.get("date", "")
                item["location"] = (payload.get("location") or {}).get("name", "")
                item["reporterName"] = payload.get("reporterName", "")
                item["folderDate"] = payload.get("_folderDate") or payload.get("date", "")
                item["hasPdf"] = bool(payload.get("_pdfPath"))
                reports.append(item)
            return json_response(start_response, {"reports": reports, "scope": context["kind"], "projectLead": context.get("name")})

        if path == "/api/reports" and method == "POST":
            data = normalize_report(read_json(environ))
            report_id = str(data.get("id") or uuid.uuid4())
            if not re.fullmatch(r"[0-9a-fA-F-]{36}", report_id):
                report_id = str(uuid.uuid4())
            now = utc_now()
            data["id"] = report_id
            title = safe_title(data)
            status = str(data.get("status") or "Entwurf")[:40]
            with db_connection() as connection:
                existing = connection.execute("SELECT created_at, payload FROM reports WHERE id = ?", (report_id,)).fetchone()
                if existing:
                    previous = json.loads(existing["payload"])
                    for metadata_key in ("_pdfPath", "_folderDate", "_submittedAt"):
                        if previous.get(metadata_key) and not data.get(metadata_key):
                            data[metadata_key] = previous[metadata_key]
                if context["kind"] == "leader":
                    if not existing:
                        return json_response(start_response, {"error": "Projektleiter können nur zugeteilte Rapporte bearbeiten"}, "403 Forbidden")
                    if previous.get("projectLead") != context["name"]:
                        return json_response(start_response, {"error": "Keine Berechtigung für diesen Rapport"}, "403 Forbidden")
                    data["projectLead"] = previous.get("projectLead")
                    data["_pdfPath"] = previous.get("_pdfPath")
                    data["_folderDate"] = previous.get("_folderDate")
                    status = previous.get("status", "Zur Kontrolle")
                    data["status"] = status
                created_at = existing["created_at"] if existing else now
                connection.execute(
                    "INSERT OR REPLACE INTO reports (id, title, status, created_at, updated_at, payload) VALUES (?, ?, ?, ?, ?, ?)",
                    (report_id, title, status, created_at, now, json.dumps(data, ensure_ascii=False)),
                )
                connection.commit()
            if data.get("_pdfPath"):
                pdf_path = DATA_DIR / data["_pdfPath"]
                pdf_path.parent.mkdir(parents=True, exist_ok=True)
                pdf_path.write_bytes(build_pdf(data))
            return json_response(start_response, {"id": report_id, "title": title, "status": status}, "201 Created")

        if path == "/api/employees" and method == "GET":
            with db_connection() as connection:
                rows = connection.execute("SELECT id, name FROM employees WHERE active = 1 ORDER BY name COLLATE NOCASE").fetchall()
            return json_response(start_response, {"employees": [dict(row) for row in rows]})

        if path == "/api/employees" and method == "POST":
            if context["kind"] != "admin":
                return json_response(start_response, {"error": "Keine Berechtigung"}, "403 Forbidden")
            data = read_json(environ)
            name = str(data.get("name") or "").strip()
            if len(name) < 2:
                raise ValueError("Bitte einen gültigen Namen eingeben")
            employee_id = str(uuid.uuid4())
            with db_connection() as connection:
                try:
                    connection.execute("INSERT INTO employees (id, name, created_at) VALUES (?, ?, ?)", (employee_id, name, utc_now()))
                    connection.commit()
                except sqlite3.IntegrityError:
                    raise ValueError("Dieser Mitarbeiter ist bereits gespeichert")
            return json_response(start_response, {"employee": {"id": employee_id, "name": name}}, "201 Created")

        employee_match = re.fullmatch(r"/api/employees/([0-9a-fA-F-]{36})", path)
        if employee_match and method == "DELETE":
            if context["kind"] != "admin":
                return json_response(start_response, {"error": "Keine Berechtigung"}, "403 Forbidden")
            with db_connection() as connection:
                connection.execute("UPDATE employees SET active = 0 WHERE id = ?", (employee_match.group(1),))
                connection.commit()
            return json_response(start_response, {"deleted": True})

        if path == "/api/hour-summary" and method == "GET":
            query = parse_qs(environ.get("QUERY_STRING", ""))
            wanted_date = (query.get("date") or [""])[0]
            totals = {}
            with db_connection() as connection:
                rows = connection.execute("SELECT status, payload FROM reports").fetchall()
            for row in rows:
                report = json.loads(row["payload"])
                if context["kind"] == "leader" and report.get("projectLead") != context["name"]:
                    continue
                if wanted_date and report.get("date") != wanted_date:
                    continue
                for area, entries in (report.get("work") or {}).items():
                    for entry in entries:
                        name = str(entry.get("name") or "").strip()
                        hours = float(entry.get("hours") or 0)
                        if not name or not hours:
                            continue
                        key = (report.get("date", ""), name)
                        current = totals.setdefault(key, {"date": key[0], "name": name, "hours": 0, "reports": 0, "areas": set()})
                        current["hours"] += hours
                        current["reports"] += 1
                        current["areas"].add(area)
            summary = []
            for item in totals.values():
                item["hours"] = round(item["hours"], 2)
                item["areas"] = sorted(item["areas"])
                summary.append(item)
            summary.sort(key=lambda item: (item["date"], item["name"]), reverse=True)
            return json_response(start_response, {"summary": summary})

        match = re.fullmatch(r"/api/reports/([0-9a-fA-F-]{36})(/(?:excel|pdf))?", path)
        if match:
            report_id, file_suffix = match.group(1), match.group(2)
            with db_connection() as connection:
                row = connection.execute("SELECT * FROM reports WHERE id = ?", (report_id,)).fetchone()
            if not row:
                return json_response(start_response, {"error": "Rapport nicht gefunden"}, "404 Not Found")
            report = json.loads(row["payload"])
            if context["kind"] == "leader" and report.get("projectLead") != context["name"]:
                return json_response(start_response, {"error": "Keine Berechtigung für diesen Rapport"}, "403 Forbidden")
            if method == "GET" and file_suffix == "/excel":
                binary = build_excel(report)
                filename = re.sub(r"[^A-Za-z0-9._-]+", "_", row["title"]) + ".xlsx"
                start_response("200 OK", [
                    ("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                    ("Content-Disposition", f'attachment; filename="{filename}"'),
                    ("Content-Length", str(len(binary))),
                ])
                return [binary]
            if method == "GET" and file_suffix == "/pdf":
                relative = report.get("_pdfPath")
                binary = (DATA_DIR / relative).read_bytes() if relative and (DATA_DIR / relative).exists() else build_pdf(report)
                filename = re.sub(r"[^A-Za-z0-9._-]+", "_", row["title"]) + ".pdf"
                start_response("200 OK", [("Content-Type", "application/pdf"), ("Content-Disposition", f'attachment; filename="{filename}"'), ("Content-Length", str(len(binary)))])
                return [binary]
            if method == "GET" and not file_suffix:
                return json_response(start_response, {"report": report})
            if method == "DELETE" and not file_suffix:
                if context["kind"] != "admin":
                    return json_response(start_response, {"error": "Keine Berechtigung"}, "403 Forbidden")
                with db_connection() as connection:
                    connection.execute("DELETE FROM reports WHERE id = ?", (report_id,))
                    connection.commit()
                return json_response(start_response, {"deleted": True})

        workflow_match = re.fullmatch(r"/api/reports/([0-9a-fA-F-]{36})/workflow", path)
        if workflow_match and method == "POST":
            data = read_json(environ)
            requested = str(data.get("status") or "")
            allowed = {"Entwurf", "Zur Kontrolle", "Frontoffice", "Erledigt"}
            if requested not in allowed:
                raise ValueError("Ungültiger Workflow-Status")
            with db_connection() as connection:
                row = connection.execute("SELECT payload FROM reports WHERE id = ?", (workflow_match.group(1),)).fetchone()
                if not row:
                    return json_response(start_response, {"error": "Rapport nicht gefunden"}, "404 Not Found")
                payload = json.loads(row["payload"])
                if context["kind"] == "leader":
                    if payload.get("projectLead") != context["name"] or requested not in {"Zur Kontrolle", "Frontoffice"}:
                        return json_response(start_response, {"error": "Keine Berechtigung für diesen Status"}, "403 Forbidden")
                if requested == "Zur Kontrolle":
                    folder_date = datetime.now(ZoneInfo("Europe/Zurich")).date().isoformat()
                    payload["_folderDate"] = folder_date
                    payload["_submittedAt"] = utc_now()
                payload["status"] = requested
                if requested == "Zur Kontrolle":
                    payload["_pdfPath"] = save_pdf_to_folder(payload)
                connection.execute("UPDATE reports SET status = ?, payload = ?, updated_at = ? WHERE id = ?", (requested, json.dumps(payload, ensure_ascii=False), utc_now(), workflow_match.group(1)))
                connection.commit()
            mail_sent, mail_detail = (False, None)
            if requested == "Zur Kontrolle":
                mail_sent, mail_detail = send_project_leader_email(payload, environ)
            return json_response(start_response, {"id": workflow_match.group(1), "status": requested, "pdfCreated": bool(payload.get("_pdfPath")), "mailSent": mail_sent, "mailDetail": mail_detail})

        if path.startswith("/api/"):
            return json_response(start_response, {"error": "Nicht gefunden"}, "404 Not Found")
        return serve_static(start_response, path)
    except (ValueError, json.JSONDecodeError) as error:
        return json_response(start_response, {"error": str(error)}, "400 Bad Request")
    except Exception as error:
        print(f"Request error: {error}", flush=True)
        return json_response(start_response, {"error": "Interner Fehler"}, "500 Internal Server Error")


if __name__ == "__main__":
    from wsgiref.simple_server import make_server

    port = int(os.environ.get("PORT", "8080"))
    print(f"Tagesrapport läuft auf http://localhost:{port}")
    make_server("0.0.0.0", port, application).serve_forever()
